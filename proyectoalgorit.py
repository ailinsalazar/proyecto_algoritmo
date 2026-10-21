import os  # permite al programa interactuar con el sistema operativo
import smtplib  # conecta nuestro correo electrónico
from email.mime.multipart import MIMEMultipart  # estudia la estructura del correo
from email.mime.text import MIMEText  # redacta el mensaje de nuestro correo
from email.mime.base import MIMEBase  # adjunta los archivos al correo
from email import encoders  # protege el archivo para que no se dañe
import tkinter as tk 
from tkinter import messagebox, simpledialog 
import openpyxl  # lee datos de excel
from docx import Document

ARCHIVO_EXCEL = "Inventario.xlsx"

# Reemplaza con tus credenciales de Gmail o el servicio SMTP que utilices
REMITENTE_EMAIL = "emilianoesquite19@gmail.com"
REMITENTE_PASSWORD = "umcfrxcwrsdxfxpy"  # Contraseña de aplicación de Gmail
CORREO_COORDINADOR = "namjoon8002@gmail.com"  # Correo para enviar la copia de seguridad

def inicializar_excel(nombre_archivo=ARCHIVO_EXCEL):
    """Crea el archivo Excel con sus 3 hojas si aún no existe."""
    if not os.path.exists(nombre_archivo):  # verifica si el archivo excel existe
        wb = openpyxl.Workbook()  # crea un libro excel

        # Hoja 1: productos (nombre, precio, existencia)
        ws_prod = wb.active
        ws_prod.title = "productos"
        ws_prod.append(["ID", "Nombre", "Precio", "Existencia"])

        # Hoja 2: clientes (nombre, nit, dirección, email)
        ws_cli = wb.create_sheet(title="clientes")
        ws_cli.append(["ID", "Nombre", "NIT", "Dirección", "Email"])

        # Hoja 3: pedidos (nombre cliente, nombre producto, cantidad producto, valor pedido)
        ws_ped = wb.create_sheet(title="pedidos")
        ws_ped.append(["ID_Pedido", "Nombre Cliente", "Nombre Producto", "Cantidad Producto", "Valor Pedido"])

        wb.save(nombre_archivo)


def mostrar_frame(frame):
    for f in frames.values():
        f.pack_forget()
    frame.pack(fill="both", expand=True)


def obtener_siguiente_id(ws):
    max_id = 0
    for row in ws.iter_rows(min_row=2, values_only=True):  # revisa las filas de excel desde la fila 2
        if row[0] and isinstance(row[0], int):  # verifica que tenga un valor válido
            if row[0] > max_id:
                max_id = row[0]  # compara los IDs existentes
    return max_id + 1


def guardar_producto():
    nombre = entry_prod_nombre.get().strip()
    precio_str = entry_prod_precio.get().strip()
    stock_str = entry_prod_stock.get().strip()

    if not nombre or not precio_str or not stock_str:
        messagebox.showwarning("Campos incompletos", "Por favor completa todos los campos.")
        return

    try:
        precio = float(precio_str)
        stock = int(stock_str)
        if precio < 0 or stock < 0:
            raise ValueError
    except ValueError:
        messagebox.showerror("Datos inválidos", "El precio y existencia deben ser números mayor o igual a 0.")
        return

    try:
        wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
        ws = wb["productos"]
        nuevo_id = obtener_siguiente_id(ws)
        ws.append([nuevo_id, nombre, precio, stock])
        wb.save(ARCHIVO_EXCEL)

        messagebox.showinfo("Éxito", f"Producto '{nombre}' guardado con ID #{nuevo_id}.")
        entry_prod_nombre.delete(0, tk.END)
        entry_prod_precio.delete(0, tk.END)
        entry_prod_stock.delete(0, tk.END)
    except Exception as e:
        messagebox.showerror("Error", f"Ocurrió un problema: {e}")


def listar_productos():
    try:
        wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
        ws = wb["productos"]
        filas = list(ws.iter_rows(values_only=True))

        if len(filas) <= 1:
            messagebox.showinfo("Productos", "No hay productos registrados aún.")
            return

        info = "ID | Nombre | Precio | Existencia\n" + "-" * 45 + "\n"
        for row in filas[1:]:
            info += f"#{row[0]} | {row[1]} | Q.{row[2]:.2f} | {row[3]} unidades\n"
        messagebox.showinfo("Lista de Productos", info)
    except Exception as e:
        messagebox.showerror("Error", f"No se pudo leer el archivo: {e}")


def editar_producto():
    prod_id = simpledialog.askinteger("Editar Producto", "Ingrese el ID del producto a editar:")
    if not prod_id:
        return

    try:
        wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
        ws = wb["productos"]
        encontrado = False

        for row in ws.iter_rows(min_row=2):  # despliega diálogos emergentes con el valor actual
            if row[0].value == prod_id:
                nuevo_nombre = simpledialog.askstring("Editar Nombre", "Nuevo Nombre:", initialvalue=row[1].value)
                nuevo_precio = simpledialog.askfloat("Editar Precio", "Nuevo Precio (Q):", initialvalue=row[2].value)
                nuevo_stock = simpledialog.askinteger("Editar Existencia", "Nueva Existencia:", initialvalue=row[3].value)

                if nuevo_nombre and nuevo_precio is not None and nuevo_stock is not None:
                    row[1].value = nuevo_nombre
                    row[2].value = nuevo_precio
                    row[3].value = nuevo_stock
                    wb.save(ARCHIVO_EXCEL)
                    messagebox.showinfo("Éxito", "Producto actualizado correctamente.")
                encontrado = True
                break

        if not encontrado:
            messagebox.showerror("Error", "Producto no encontrado.")
    except Exception as e:
        messagebox.showerror("Error", f"No se pudo editar el producto: {e}")


def eliminar_producto():
    prod_id = simpledialog.askinteger("Eliminar Producto", "Ingrese el ID del producto a eliminar:")
    if not prod_id:
        return

    try:
        wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
        ws = wb["productos"]
        encontrado = False

        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if row[0] == prod_id:
                ws.delete_rows(idx)
                wb.save(ARCHIVO_EXCEL)
                messagebox.showinfo("Éxito", "Producto eliminado correctamente.")
                encontrado = True
                break

        if not encontrado:
            messagebox.showerror("Error", "Producto no encontrado.")
    except Exception as e:
        messagebox.showerror("Error", f"No se pudo eliminar el producto: {e}")


def enviar_cotizacion_correo():
    cliente_nom = simpledialog.askstring("Cotización", "Nombre del cliente:")
    cliente_email = simpledialog.askstring("Cotización", "Correo del cliente:")
    prod_nom = simpledialog.askstring("Cotización", "Nombre del producto a cotizar:")

    if not cliente_nom or not cliente_email or not prod_nom:
        messagebox.showwarning("Atención", "Todos los datos son requeridos para la cotización.")
        return

    # Buscar precio del producto en Excel
    precio_prod = None
    try:
        wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
        ws = wb["productos"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[1] and str(row[1]).lower() == prod_nom.lower():
                precio_prod = row[2]
                break
    except Exception as e:
        messagebox.showerror("Error", f"Error al consultar base de datos: {e}")
        return

    if precio_prod is None:
        messagebox.showerror("Error", f"El producto '{prod_nom}' no fue encontrado en el inventario.")
        return

    # Generar documento Word 'cotización.docx'
    nombre_archivo_doc = "cotización.docx"
    try:
        doc = Document()
        doc.add_heading("COTIZACIÓN DE PRODUCTO", level=1)
        
        texto_plantilla = f"Estimado {cliente_nom}\nEl precio de nuestro producto {prod_nom} es Q.{precio_prod:.2f}"
        doc.add_paragraph(texto_plantilla)
        doc.save(nombre_archivo_doc)
    except Exception as e:
        messagebox.showerror("Error Word", f"No se pudo generar el documento Word: {e}")
        return

    # Enviar por correo SMTP
    try:
        msg = MIMEMultipart()
        msg['From'] = REMITENTE_EMAIL
        msg['To'] = cliente_email
        msg['Subject'] = f"Cotización de Producto - {prod_nom}"

        cuerpo = f"Hola {cliente_nom},\n\nAdjunto encontrarás la cotización formal solicitada.\n\nSaludos cordiales."
        msg.attach(MIMEText(cuerpo, 'plain'))

        # Adjuntar Word
        with open(nombre_archivo_doc, "rb") as attachment:
            part = MIMEBase("application", "octet-stream")
            part.set_payload(attachment.read())
            encoders.encode_base64(part)
            part.add_header("Content-Disposition", f"attachment; filename= {nombre_archivo_doc}")
            msg.attach(part)

        server = smtplib.SMTP('smtp.gmail.com', 587)
        server.starttls()
        server.login(REMITENTE_EMAIL, REMITENTE_PASSWORD)
        server.send_message(msg)
        server.quit()

        messagebox.showinfo("Correo Enviado", f"Cotización enviada exitosamente a {cliente_email}.")
    except Exception as e:
        messagebox.showerror("Error SMTP", f"No se pudo enviar el correo: {e}\n(Verifique sus credenciales SMTP en el código)")


def guardar_cliente():
    nombre = entry_cli_nombre.get().strip()
    nit = entry_cli_nit.get().strip()
    direccion = entry_cli_dir.get().strip()
    email = entry_cli_email.get().strip()

    if not nombre or not nit or not direccion or not email:
        messagebox.showwarning("Campos incompletos", "Por favor completa todos los campos del cliente.")
        return

    try:
        wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
        ws = wb["clientes"]
        nuevo_id = obtener_siguiente_id(ws)
        ws.append([nuevo_id, nombre, nit, direccion, email])
        wb.save(ARCHIVO_EXCEL)

        messagebox.showinfo("Éxito", f"Cliente '{nombre}' guardado con ID #{nuevo_id}.")
        entry_cli_nombre.delete(0, tk.END)
        entry_cli_nit.delete(0, tk.END)
        entry_cli_dir.delete(0, tk.END)
        entry_cli_email.delete(0, tk.END)
    except Exception as e:
        messagebox.showerror("Error", f"Ocurrió un problema al guardar cliente: {e}")


def listar_clientes():
    try:
        wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
        ws = wb["clientes"]
        filas = list(ws.iter_rows(values_only=True))

        if len(filas) <= 1:
            messagebox.showinfo("Clientes", "No hay clientes registrados aún.")
            return

        info = "ID | Nombre | NIT | Dirección | Email\n" + "-" * 55 + "\n"
        for row in filas[1:]:
            info += f"#{row[0]} | {row[1]} | NIT: {row[2]} | {row[3]} | {row[4]}\n"
        messagebox.showinfo("Lista de Clientes", info)
    except Exception as e:
        messagebox.showerror("Error", f"No se pudo leer el archivo: {e}")


def editar_cliente():
    cli_id = simpledialog.askinteger("Editar Cliente", "Ingrese el ID del cliente a editar:")
    if not cli_id:
        return

    try:
        wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
        ws = wb["clientes"]
        encontrado = False

        for row in ws.iter_rows(min_row=2):
            if row[0].value == cli_id:
                nuevo_nom = simpledialog.askstring("Editar Nombre", "Nuevo Nombre:", initialvalue=row[1].value)
                nuevo_nit = simpledialog.askstring("Editar NIT", "Nuevo NIT:", initialvalue=row[2].value)
                nueva_dir = simpledialog.askstring("Editar Dirección", "Nueva Dirección:", initialvalue=row[3].value)
                nuevo_email = simpledialog.askstring("Editar Email", "Nuevo Email:", initialvalue=row[4].value)

                if nuevo_nom and nuevo_nit and nueva_dir and nuevo_email:
                    row[1].value = nuevo_nom
                    row[2].value = nuevo_nit
                    row[3].value = nueva_dir
                    row[4].value = nuevo_email
                    wb.save(ARCHIVO_EXCEL)
                    messagebox.showinfo("Éxito", "Cliente actualizado correctamente.")
                encontrado = True
                break

        if not encontrado:
            messagebox.showerror("Error", "Cliente no encontrado.")
    except Exception as e:
        messagebox.showerror("Error", f"No se pudo editar el cliente: {e}")


def eliminar_cliente():
    cli_id = simpledialog.askinteger("Eliminar Cliente", "Ingrese el ID del cliente a eliminar:")
    if not cli_id:
        return

    try:
        wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
        ws = wb["clientes"]
        encontrado = False

        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if row[0] == cli_id:
                ws.delete_rows(idx)
                wb.save(ARCHIVO_EXCEL)
                messagebox.showinfo("Éxito", "Cliente eliminado correctamente.")
                encontrado = True
                break

        if not encontrado:
            messagebox.showerror("Error", "Cliente no encontrado.")
    except Exception as e:
        messagebox.showerror("Error", f"No se pudo eliminar el cliente: {e}")


def guardar_pedido():
    cliente_nom = entry_ped_cliente.get().strip()
    producto_nom = entry_ped_producto.get().strip()
    cant_str = entry_ped_cant.get().strip()

    if not cliente_nom or not producto_nom or not cant_str:
        messagebox.showwarning("Campos incompletos", "Por favor completa todos los campos del pedido.")
        return

    try:
        cant = int(cant_str)
        if cant <= 0:
            raise ValueError
    except ValueError:
        messagebox.showerror("Cantidad inválida", "La cantidad debe ser un entero positivo mayor a 0.")
        return

    # Buscar el precio del producto y descontar stock
    wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
    ws_prod = wb["productos"]
    precio_unitario = None
    prod_row_idx = None
    stock_actual = 0

    for idx, row in enumerate(ws_prod.iter_rows(min_row=2), start=2):
        if row[1].value and str(row[1].value).lower() == producto_nom.lower():
            precio_unitario = float(row[2].value)
            stock_actual = int(row[3].value)
            prod_row_idx = idx
            break

    if precio_unitario is None:
        messagebox.showerror("Error", f"El producto '{producto_nom}' no existe en el inventario.")
        return

    if cant > stock_actual:
        messagebox.showerror("Stock insuficiente", f"Solo quedan {stock_actual} unidades en inventario.")
        return

    # Actualizar stock del producto
    ws_prod.cell(row=prod_row_idx, column=4, value=stock_actual - cant)

    # Registrar el pedido
    ws_ped = wb["pedidos"]
    nuevo_id_ped = obtener_siguiente_id(ws_ped)
    total_pedido = precio_unitario * cant

    ws_ped.append([nuevo_id_ped, cliente_nom, producto_nom, cant, total_pedido])
    wb.save(ARCHIVO_EXCEL)

    messagebox.showinfo("Éxito", f"Pedido #{nuevo_id_ped} registrado exitosamente. Total: Q.{total_pedido:.2f}")
    entry_ped_cliente.delete(0, tk.END)
    entry_ped_producto.delete(0, tk.END)
    entry_ped_cant.delete(0, tk.END)


def listar_pedidos():
    try:
        wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
        ws = wb["pedidos"]
        filas = list(ws.iter_rows(values_only=True))

        if len(filas) <= 1:
            messagebox.showinfo("Pedidos", "No hay pedidos registrados aún.")
            return

        info = "ID Pedido | Cliente | Producto | Cantidad | Total\n" + "-" * 55 + "\n"
        for row in filas[1:]:
            info += f"#{row[0]} | {row[1]} | {row[2]} | {row[3]} unid. | Q.{row[4]:.2f}\n"
        messagebox.showinfo("Lista de Pedidos", info)
    except Exception as e:
        messagebox.showerror("Error", f"No se pudo leer el archivo: {e}")


def eliminar_pedido():
    ped_id = simpledialog.askinteger("Eliminar Pedido", "Ingrese el ID del pedido a eliminar:")
    if not ped_id:
        return

    try:
        wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
        ws = wb["pedidos"]
        encontrado = False

        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if row[0] == ped_id:
                ws.delete_rows(idx)
                wb.save(ARCHIVO_EXCEL)
                messagebox.showinfo("Éxito", "Pedido eliminado correctamente.")
                encontrado = True
                break

        if not encontrado:
            messagebox.showerror("Error", "Pedido no encontrado.")
    except Exception as e:
        messagebox.showerror("Error", f"No se pudo eliminar el pedido: {e}")


def informe_ventas_cliente():
    try:
        wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
        ws = wb["pedidos"]
        totales_cliente = {}

        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[1] and row[4] is not None:
                cliente = row[1]
                total = float(row[4])
                totales_cliente[cliente] = totales_cliente.get(cliente, 0.0) + total

        if not totales_cliente:
            messagebox.showinfo("Informes", "No hay ventas registradas para generar el informe.")
            return

        info = "--- TOTAL DE VENTAS POR CLIENTE ---\n\n"
        for cliente, total in totales_cliente.items():
            info += f"• Cliente: {cliente} -> Total Comprado: Q.{total:.2f}\n"

        messagebox.showinfo("Informe - Ventas por Cliente", info)
    except Exception as e:
        messagebox.showerror("Error", f"Error al generar informe: {e}")


def informe_ventas_producto():
    try:
        wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
        ws = wb["pedidos"]
        totales_producto = {}

        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[2] and row[4] is not None:
                producto = row[2]
                total = float(row[4])
                totales_producto[producto] = totales_producto.get(producto, 0.0) + total

        if not totales_producto:
            messagebox.showinfo("Informes", "No hay ventas registradas para generar el informe.")
            return

        info = "--- TOTAL DE VENTAS POR PRODUCTO ---\n\n"
        for producto, total in totales_producto.items():
            info += f"• Producto: {producto} -> Total Vendido: Q.{total:.2f}\n"

        messagebox.showinfo("Informe - Ventas por Producto", info)
    except Exception as e:
        messagebox.showerror("Error", f"Error al generar informe: {e}")


def crear_copia_seguridad():
    if not os.path.exists(ARCHIVO_EXCEL):
        messagebox.showerror("Error", "No existe el archivo de datos para respaldar.")
        return

    try:
        msg = MIMEMultipart()
        msg['From'] = REMITENTE_EMAIL
        msg['To'] = CORREO_COORDINADOR
        msg['Subject'] = "Copia de Seguridad - Inventario.xlsx"

        cuerpo = "Estimado Coordinador,\n\nAdjunto se encuentra la copia de seguridad del archivo Inventario.xlsx."
        msg.attach(MIMEText(cuerpo, 'plain'))

        # Adjuntar Excel completo
        with open(ARCHIVO_EXCEL, "rb") as attachment:
            part = MIMEBase("application", "octet-stream")
            part.set_payload(attachment.read())
            encoders.encode_base64(part)
            part.add_header("Content-Disposition", f"attachment; filename= {ARCHIVO_EXCEL}")
            msg.attach(part)

        server = smtplib.SMTP('smtp.gmail.com', 587)
        server.starttls()
        server.login(REMITENTE_EMAIL, REMITENTE_PASSWORD)
        server.send_message(msg)
        server.quit()

        messagebox.showinfo("Copia de Seguridad", f"Copia de seguridad enviada con éxito a {CORREO_COORDINADOR}.")
    except Exception as e:
        messagebox.showerror("Error SMTP", f"No se pudo enviar la copia de seguridad: {e}\n(Verifique credenciales SMTP en el código)")


# --- EJECUCIÓN DE INTERFAZ GRÁFICA ---
if __name__ == "__main__":
    inicializar_excel()

    ventana = tk.Tk()
    ventana.title("Sistema de Inventario")
    ventana.geometry("750x650")
    ventana.configure(bg="#E1DCCA")

    frames = {}

    frame_menu = tk.Frame(ventana, bg="#E1DCCA")
    frames["menu"] = frame_menu

    titulo = tk.Label(
        frame_menu,
        text="SISTEMA DE INVENTARIO",
        font=("Georgia", 28, "bold"),
        fg="maroon",
        bg="#E1DCCA",
    )
    titulo.pack(pady=20)

    btn_opts = [
        ("Productos", "productos"),
        ("Clientes", "clientes"),
        ("Pedidos", "pedidos"),
        ("Informes", "informes"),
        ("Copia de seguridad", "respaldo"),
    ]

    for texto, clave in btn_opts:
        tk.Button(
            frame_menu,
            text=texto,
            font=("Georgia", 16),
            width=25,
            bg="#003049",
            fg="#FABD32",
            command=lambda c=clave: mostrar_frame(frames[c]),
        ).pack(pady=6)

    frame_menu.pack(fill="both", expand=True)

    frame_productos = tk.Frame(ventana, bg="#E1DCCA")
    frames["productos"] = frame_productos

    tk.Label(frame_productos, text="Gestión de Productos", font=("Georgia", 22, "bold"), fg="maroon", bg="#E1DCCA").pack(pady=10)

    frame_form_prod = tk.LabelFrame(frame_productos, text=" Agregar Nuevo Producto ", font=("Georgia", 12, "bold"), bg="#E1DCCA", fg="#003049", padx=10, pady=10)
    frame_form_prod.pack(pady=5, padx=20)

    tk.Label(frame_form_prod, text="Nombre:", font=("Georgia", 10), bg="#E1DCCA").grid(row=0, column=0, sticky="e", pady=3, padx=5)
    entry_prod_nombre = tk.Entry(frame_form_prod, font=("Georgia", 10), width=25)
    entry_prod_nombre.grid(row=0, column=1, pady=3, padx=5)

    tk.Label(frame_form_prod, text="Precio (Q):", font=("Georgia", 10), bg="#E1DCCA").grid(row=1, column=0, sticky="e", pady=3, padx=5)
    entry_prod_precio = tk.Entry(frame_form_prod, font=("Georgia", 10), width=25)
    entry_prod_precio.grid(row=1, column=1, pady=3, padx=5)

    tk.Label(frame_form_prod, text="Existencia inicial:", font=("Georgia", 10), bg="#E1DCCA").grid(row=2, column=0, sticky="e", pady=3, padx=5)
    entry_prod_stock = tk.Entry(frame_form_prod, font=("Georgia", 10), width=25)
    entry_prod_stock.grid(row=2, column=1, pady=3, padx=5)

    btn_guardar_prod = tk.Button(frame_form_prod, text="Guardar Producto", font=("Georgia", 10, "bold"), bg="#003049", fg="#FABD32", command=guardar_producto)
    btn_guardar_prod.grid(row=3, column=0, columnspan=2, pady=8)

    frame_acciones_prod = tk.Frame(frame_productos, bg="#E1DCCA")
    frame_acciones_prod.pack(pady=5)

    tk.Button(frame_acciones_prod, text="Editar producto", font=("Georgia", 11), width=22, bg="#003049", fg="#FABD32", command=editar_producto).pack(pady=3)
    tk.Button(frame_acciones_prod, text="Eliminar producto", font=("Georgia", 11), width=22, bg="#003049", fg="#FABD32", command=eliminar_producto).pack(pady=3)
    tk.Button(frame_acciones_prod, text="Listar productos", font=("Georgia", 11), width=22, bg="#003049", fg="#FABD32", command=listar_productos).pack(pady=3)
    tk.Button(frame_acciones_prod, text="Enviar cotización por correo", font=("Georgia", 11), width=22, bg="#003049", fg="#FABD32", command=enviar_cotizacion_correo).pack(pady=3)

    tk.Button(frame_productos, text="Volver al menú", font=("Georgia", 11, "bold"), width=22, bg="#F77F00", fg="#003049", command=lambda: mostrar_frame(frames["menu"])).pack(pady=10)

    frame_clientes = tk.Frame(ventana, bg="#E1DCCA")
    frames["clientes"] = frame_clientes

    tk.Label(frame_clientes, text="Gestión de Clientes", font=("Georgia", 22, "bold"), fg="maroon", bg="#E1DCCA").pack(pady=10)

    frame_form_cli = tk.LabelFrame(frame_clientes, text=" Agregar Nuevo Cliente ", font=("Georgia", 12, "bold"), bg="#E1DCCA", fg="#003049", padx=10, pady=10)
    frame_form_cli.pack(pady=5, padx=20)

    tk.Label(frame_form_cli, text="Nombre:", font=("Georgia", 10), bg="#E1DCCA").grid(row=0, column=0, sticky="e", pady=3, padx=5)
    entry_cli_nombre = tk.Entry(frame_form_cli, font=("Georgia", 10), width=25)
    entry_cli_nombre.grid(row=0, column=1, pady=3, padx=5)

    tk.Label(frame_form_cli, text="NIT:", font=("Georgia", 10), bg="#E1DCCA").grid(row=1, column=0, sticky="e", pady=3, padx=5)
    entry_cli_nit = tk.Entry(frame_form_cli, font=("Georgia", 10), width=25)
    entry_cli_nit.grid(row=1, column=1, pady=3, padx=5)

    tk.Label(frame_form_cli, text="Dirección:", font=("Georgia", 10), bg="#E1DCCA").grid(row=2, column=0, sticky="e", pady=3, padx=5)
    entry_cli_dir = tk.Entry(frame_form_cli, font=("Georgia", 10), width=25)
    entry_cli_dir.grid(row=2, column=1, pady=3, padx=5)

    tk.Label(frame_form_cli, text="Email:", font=("Georgia", 10), bg="#E1DCCA").grid(row=3, column=0, sticky="e", pady=3, padx=5)
    entry_cli_email = tk.Entry(frame_form_cli, font=("Georgia", 10), width=25)
    entry_cli_email.grid(row=3, column=1, pady=3, padx=5)

    tk.Button(frame_form_cli, text="Guardar Cliente", font=("Georgia", 10, "bold"), bg="#003049", fg="#FABD32", command=guardar_cliente).grid(row=4, column=0, columnspan=2, pady=8)

    frame_acciones_cli = tk.Frame(frame_clientes, bg="#E1DCCA")
    frame_acciones_cli.pack(pady=5)

    tk.Button(frame_acciones_cli, text="Editar cliente", font=("Georgia", 11), width=22, bg="#003049", fg="#FABD32", command=editar_cliente).pack(pady=3)
    tk.Button(frame_acciones_cli, text="Eliminar cliente", font=("Georgia", 11), width=22, bg="#003049", fg="#FABD32", command=eliminar_cliente).pack(pady=3)
    tk.Button(frame_acciones_cli, text="Listar clientes", font=("Georgia", 11), width=22, bg="#003049", fg="#FABD32", command=listar_clientes).pack(pady=3)

    tk.Button(frame_clientes, text="Volver al menú", font=("Georgia", 11, "bold"), width=22, bg="#F77F00", fg="#003049", command=lambda: mostrar_frame(frames["menu"])).pack(pady=10)

    frame_pedidos = tk.Frame(ventana, bg="#E1DCCA")
    frames["pedidos"] = frame_pedidos

    tk.Label(frame_pedidos, text="Gestión de Pedidos", font=("Georgia", 22, "bold"), fg="maroon", bg="#E1DCCA").pack(pady=10)

    frame_form_ped = tk.LabelFrame(frame_pedidos, text=" Agregar Nuevo Pedido ", font=("Georgia", 12, "bold"), bg="#E1DCCA", fg="#003049", padx=10, pady=10)
    frame_form_ped.pack(pady=5, padx=20)

    tk.Label(frame_form_ped, text="Nombre Cliente:", font=("Georgia", 10), bg="#E1DCCA").grid(row=0, column=0, sticky="e", pady=3, padx=5)
    entry_ped_cliente = tk.Entry(frame_form_ped, font=("Georgia", 10), width=25)
    entry_ped_cliente.grid(row=0, column=1, pady=3, padx=5)

    tk.Label(frame_form_ped, text="Nombre Producto:", font=("Georgia", 10), bg="#E1DCCA").grid(row=1, column=0, sticky="e", pady=3, padx=5)
    entry_ped_producto = tk.Entry(frame_form_ped, font=("Georgia", 10), width=25)
    entry_ped_producto.grid(row=1, column=1, pady=3, padx=5)

    tk.Label(frame_form_ped, text="Cantidad:", font=("Georgia", 10), bg="#E1DCCA").grid(row=2, column=0, sticky="e", pady=3, padx=5)
    entry_ped_cant = tk.Entry(frame_form_ped, font=("Georgia", 10), width=25)
    entry_ped_cant.grid(row=2, column=1, pady=3, padx=5)

    tk.Button(frame_form_ped, text="Guardar Pedido", font=("Georgia", 10, "bold"), bg="#003049", fg="#FABD32", command=guardar_pedido).grid(row=3, column=0, columnspan=2, pady=8)

    frame_acciones_ped = tk.Frame(frame_pedidos, bg="#E1DCCA")
    frame_acciones_ped.pack(pady=5)

    tk.Button(frame_acciones_ped, text="Eliminar pedido", font=("Georgia", 11), width=22, bg="#003049", fg="#FABD32", command=eliminar_pedido).pack(pady=3)
    tk.Button(frame_acciones_ped, text="Listar pedidos", font=("Georgia", 11), width=22, bg="#003049", fg="#FABD32", command=listar_pedidos).pack(pady=3)

    tk.Button(frame_pedidos, text="Volver al menú", font=("Georgia", 11, "bold"), width=22, bg="#F77F00", fg="#003049", command=lambda: mostrar_frame(frames["menu"])).pack(pady=10)

    frame_informes = tk.Frame(ventana, bg="#E1DCCA")
    frames["informes"] = frame_informes

    tk.Label(frame_informes, text="Informes y Reportes", font=("Georgia", 22, "bold"), fg="maroon", bg="#E1DCCA").pack(pady=20)

    frame_acciones_inf = tk.Frame(frame_informes, bg="#E1DCCA")
    frame_acciones_inf.pack(pady=10)

    tk.Button(frame_acciones_inf, text="Total de venta por cliente", font=("Georgia", 12), width=25, bg="#003049", fg="#FABD32", command=informe_ventas_cliente).pack(pady=8)
    tk.Button(frame_acciones_inf, text="Total de ventas por producto", font=("Georgia", 12), width=25, bg="#003049", fg="#FABD32", command=informe_ventas_producto).pack(pady=8)

    tk.Button(frame_informes, text="Volver al menú", font=("Georgia", 11, "bold"), width=22, bg="#F77F00", fg="#003049", command=lambda: mostrar_frame(frames["menu"])).pack(pady=20)

    frame_respaldo = tk.Frame(ventana, bg="#E1DCCA")
    frames["respaldo"] = frame_respaldo

    tk.Label(frame_respaldo, text="Copia de Seguridad", font=("Georgia", 22, "bold"), fg="maroon", bg="#E1DCCA").pack(pady=20)

    tk.Label(frame_respaldo, text="Envía una copia del archivo 'Inventario.xlsx'\nal coordinador del grupo por correo.", font=("Georgia", 11), bg="#E1DCCA").pack(pady=10)

    tk.Button(frame_respaldo, text="Crear copia de seguridad", font=("Georgia", 12, "bold"), width=25, bg="#003049", fg="#FABD32", command=crear_copia_seguridad).pack(pady=15)

    tk.Button(frame_respaldo, text="Volver al menú", font=("Georgia", 11, "bold"), width=22, bg="#F77F00", fg="#003049", command=lambda: mostrar_frame(frames["menu"])).pack(pady=20)

    ventana.mainloop()